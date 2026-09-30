import json
import runpy
import sys
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from PIL import Image, ImageFont
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Inches

from src import generate_capstone_decks as decks

CANONICAL_VOLUME = decks.VOLUME
MEASURED_FONT = decks._font


class DeckLayoutTests(unittest.TestCase):
    def setUp(self):
        self.directory = tempfile.TemporaryDirectory(dir=decks.ROOT)
        self.addCleanup(self.directory.cleanup)
        self.output = Path(self.directory.name)
        output_patch = patch.object(decks, "OUTPUT", self.output)
        output_patch.start()
        self.addCleanup(output_patch.stop)
        (self.output / "volume_forecast_experiment.json").write_text("{}", encoding="utf-8")
        volume_patch = patch.object(decks, "VOLUME", self.output)
        volume_patch.start()
        self.addCleanup(volume_patch.stop)
        # Isolated geometry fixtures use Pillow's bundled font, not system fonts/matplotlib.
        font_patch = patch.object(decks, "_font", side_effect=self.fixture_font)
        font_patch.start()
        self.addCleanup(font_patch.stop)

    @staticmethod
    def fixture_font(size, bold=False):
        font = ImageFont.load_default(size=size * 4)
        return font, font.getname()[0]

    def assert_layout(self, presentation):
        self.assertGreaterEqual(len(presentation.slides), 8)
        self.assertLessEqual(len(presentation.slides), 12)
        for number, slide in enumerate(presentation.slides, 1):
            dynamic = []
            for shape in slide.shapes:
                with self.subTest(slide=number, shape=shape.name):
                    self.assertGreaterEqual(shape.left, 0)
                    self.assertGreaterEqual(shape.top, 0)
                    self.assertLessEqual(shape.left + shape.width, presentation.slide_width)
                    self.assertLessEqual(shape.top + shape.height, presentation.slide_height)
                    is_footer = shape.has_text_frame and shape.text in {
                        "Help Desk Ticket SLA and Volume Forecasting | Historical analysis",
                        decks.PRE_CORRECTION_WARNING,
                        f"{number:02d} / {len(presentation.slides):02d}",
                    }
                    if number > 1 and shape.top >= Inches(1.68) and not is_footer:
                        self.assertLessEqual(shape.top + shape.height, decks.CONTENT_BOTTOM)
                        dynamic.append(shape)
            for previous, following in zip(dynamic, dynamic[1:]):
                self.assertLessEqual(previous.top + previous.height, following.top)

    def test_both_artifact_backed_decks_preserve_content_and_fit(self):
        if not (CANONICAL_VOLUME / "volume_forecast_experiment.json").is_file():
            self.skipTest("Canonical forecast artifacts are not included in the lightweight environment")
        with patch.object(decks, "VOLUME", CANONICAL_VOLUME), patch.object(
            decks, "_font", MEASURED_FONT
        ):
            self.check_artifact_backed_decks()

    def check_artifact_backed_decks(self):
        manifest = json.loads(
            (CANONICAL_VOLUME / "volume_forecast_experiment.json").read_text(encoding="utf-8")
        )
        warning_expected = manifest.get("forecast_calendar_policy") != decks.FORECAST_CALENDAR_POLICY
        for filename, entries in decks.deck_content().items():
            with self.subTest(deck=filename):
                presentation = Presentation(decks.build_deck(filename, entries))
                self.assert_layout(presentation)
                self.assertEqual(len(presentation.slides), len(entries))
                for slide, entry in zip(presentation.slides, entries):
                    text = "\n".join(shape.text for shape in slide.shapes if shape.has_text_frame)
                    self.assertEqual(decks.PRE_CORRECTION_WARNING in text, warning_expected)
                    self.assertIn(entry[0], text)
                    for bullet in entry[1] + (entry[3] if len(entry) > 3 else []):
                        self.assertIn(bullet, text)
                    if len(entry) > 2 and entry[2]:
                        table = next(shape.table for shape in slide.shapes if shape.has_table)
                        actual = [[cell.text for cell in row.cells] for row in table.rows]
                        self.assertEqual(actual, [entry[2]["headers"]] + entry[2]["rows"])
                    pictures = [shape for shape in slide.shapes if shape.shape_type == MSO_SHAPE_TYPE.PICTURE]
                    image_path = entry[4] if len(entry) > 4 else None
                    self.assertEqual(len(pictures), int(image_path is not None))
                    if image_path:
                        with Image.open(image_path) as image:
                            self.assertAlmostEqual(pictures[0].width / pictures[0].height,
                                                   image.width / image.height, places=5)
                    for shape in slide.shapes:
                        if shape.has_text_frame:
                            for paragraph in shape.text_frame.paragraphs:
                                if paragraph.font.size:
                                    self.assertGreaterEqual(paragraph.font.size.pt, 10)

    def entries(self, entry):
        return [("Title", ["Subtitle"]), entry] + [("Body", ["Short copy"])] * 6

    def test_calendar_policy_warning_on_every_slide_of_both_decks(self):
        content = {
            "technical_fixture.pptx": self.entries(("Measures", ["Historical metrics"])),
            "business_fixture.pptx": self.entries(("Recommendation", ["Human review"])),
        }
        for policy in (None, "old_calendar_policy", decks.FORECAST_CALENDAR_POLICY):
            manifest = {} if policy is None else {"forecast_calendar_policy": policy}
            (self.output / "volume_forecast_experiment.json").write_text(
                json.dumps(manifest), encoding="utf-8"
            )
            for filename, entries in content.items():
                with self.subTest(policy=policy, deck=filename), patch.object(
                    decks, "VOLUME", self.output
                ):
                    presentation = Presentation(decks.build_deck(filename, entries))
                    self.assert_layout(presentation)
                    self.assertEqual(len(presentation.slides), len(entries))
                    for slide in presentation.slides:
                        warnings = [
                            shape for shape in slide.shapes
                            if shape.has_text_frame and shape.text == decks.PRE_CORRECTION_WARNING
                        ]
                        expected = policy != decks.FORECAST_CALENDAR_POLICY
                        self.assertEqual(len(warnings), int(expected))
                        if expected:
                            warning = warnings[0]
                            self.assertEqual(warning.top, Inches(7.08))
                            self.assertTrue(warning.text_frame.paragraphs[0].font.bold)
                            self.assertLessEqual(
                                decks._text_height([warning.text], warning.width, [10], 0, True),
                                warning.height,
                            )

    def test_tall_image_shrinks_proportionally_with_note_reserved(self):
        image_path = self.output / "tall.png"
        Image.new("RGB", (400, 1000), "white").save(image_path)
        entry = ("Process", ["Introduction"], None, ["Keep this note."], image_path)
        presentation = Presentation(decks.build_deck("tall.pptx", self.entries(entry)))
        self.assert_layout(presentation)
        picture = next(s for s in presentation.slides[1].shapes
                       if s.shape_type == MSO_SHAPE_TYPE.PICTURE)
        self.assertAlmostEqual(picture.width / picture.height, 0.4, places=5)
        self.assertLess(picture.width, Inches(11))

    def test_impossible_copy_fails_without_saving(self):
        for entry in [
            ("Body", ["Long copy " * 1000]),
            ("Body", ["Short copy"], None, ["Long note " * 500]),
            ("Body", [], {"headers": ["Header"], "rows": [["Row"]] * 30}),
        ]:
            with self.subTest(entry=entry[0]), self.assertRaisesRegex(ValueError, "slide 2"):
                decks.build_deck("impossible.pptx", self.entries(entry))
            self.assertFalse((self.output / "impossible.pptx").exists())

    def test_missing_image_is_not_silently_omitted(self):
        entry = ("Process", ["Introduction"], None, [], self.output / "missing.png")
        with self.assertRaisesRegex(FileNotFoundError, "slide 2.*missing image"):
            decks.build_deck("missing.pptx", self.entries(entry))
        self.assertFalse((self.output / "missing.pptx").exists())

    def test_missing_manifest_fails_real_generation(self):
        with patch.object(decks, "VOLUME", self.output / "absent"):
            with self.assertRaises(FileNotFoundError):
                decks.build_deck("missing_manifest.pptx", self.entries(("Body", [])))
            with self.assertRaises(FileNotFoundError):
                decks.deck_content()
        self.assertFalse((self.output / "missing_manifest.pptx").exists())

    def test_direct_script_import_uses_shared_policy_without_exporting(self):
        source = Path(decks.__file__)
        with patch.object(sys, "path", [str(source.parent)] + sys.path):
            namespace = runpy.run_path(str(source), run_name="deck_import_check")
        self.assertEqual(namespace["FORECAST_CALENDAR_POLICY"], decks.FORECAST_CALENDAR_POLICY)
        self.assertTrue(callable(namespace["build_all"]))

    def test_unreadable_image_and_title_content_fail(self):
        image_path = self.output / "narrow.png"
        Image.new("RGB", (10, 1000), "white").save(image_path)
        with self.assertRaisesRegex(ValueError, "too small"):
            decks.build_deck("narrow.pptx", self.entries(("Image", [], None, [], image_path)))
        with self.assertRaisesRegex(ValueError, "insufficient space"):
            decks.build_deck("crowded.pptx", self.entries(
                ("Image", ["Long copy " * 1000], None, ["Note"], image_path)
            ))
        entries = self.entries(("Body", []))
        entries[0] = ("Title", ["Copy"] * 30)
        with self.assertRaisesRegex(ValueError, "title slide body"):
            decks.build_deck("title.pptx", entries)

    def test_slide_count_contract(self):
        for count in (7, 13):
            with self.subTest(count=count), self.assertRaisesRegex(ValueError, "8 and 12"):
                decks.build_deck("invalid.pptx", [("Title", [])] * count)


if __name__ == "__main__":
    unittest.main()
