"""Generate the technical and business decks from the current issues.csv artifacts."""

import json
from functools import lru_cache
from pathlib import Path

import pandas as pd
from PIL import Image, ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

if __package__:
    from .issues_capstone import FORECAST_CALENDAR_POLICY
else:
    from issues_capstone import FORECAST_CALENDAR_POLICY

ROOT = Path(__file__).resolve().parents[1]
REPORTS = ROOT / "output" / "issue_helpdesk"
VOLUME = REPORTS / "volume_forecast"
OUTPUT = ROOT / "presentations"
NAVY = RGBColor(31, 55, 82)
BLUE = RGBColor(52, 112, 168)
GRAY = RGBColor(75, 84, 93)
CONTENT_BOTTOM = Inches(6.93)  # Leave a clear gap above the footer at 7.08 inches.
CONTENT_GAP = Inches(0.15)
PRE_CORRECTION_WARNING = "PRE-CORRECTION RESULTS | Comparison re-run required"


@lru_cache(maxsize=None)
def _font(size, bold=False):
    from matplotlib import font_manager

    # Use the same installed font for measurement and rendering, including fallback.
    properties = font_manager.FontProperties(
        family=["Aptos", "DejaVu Sans"], weight="bold" if bold else "normal"
    )
    path = font_manager.findfont(properties)
    return ImageFont.truetype(path, size * 4), font_manager.FontProperties(fname=path).get_name()


def _text_height(texts, width, sizes, space_after, bold=False):
    """Budget wrapped lines at unchanged font sizes, including textbox margins."""
    available = (width.inches * 72 - 14.4) * 4 / 1.05
    height = 7.2
    for text, size in zip(texts, sizes):
        font, _ = _font(size, bold)
        lines = 0
        for explicit_line in text.replace("\v", "\n").split("\n"):
            line = ""
            lines += 1
            for word in explicit_line.split():
                candidate = f"{line} {word}" if line else word
                if font.getlength(candidate) <= available:
                    line = candidate
                    continue
                if line:
                    lines += 1
                line = ""
                for char in word:
                    if font.getlength(line + char) > available:
                        lines += 1
                        line = ""
                    line += char
        height += lines * size * 1.25 + space_after
    return Pt(height)


def _bullet_sizes(bullets, small=False):
    base_size, long_size = (13, 11) if small else (18, 15)
    return [base_size if len(text) < 120 else long_size for text in bullets]


def _bullet_height(bullets, small=False):
    if not bullets:
        return 0
    return _text_height(
        [f"\u2022  {text}" for text in bullets], Inches(11.5),
        _bullet_sizes(bullets, small), 8 if small else 12,
    )


def deck_content():
    manifest = json.loads((VOLUME / "volume_forecast_experiment.json").read_text(encoding="utf-8"))
    metrics = pd.read_csv(VOLUME / "volume_forecast_rolling_metrics.csv")
    sla = pd.read_csv(REPORTS / "sla_attainment_by_priority.csv")
    quality = pd.read_csv(REPORTS / "source_quality_summary.csv").set_index("quality_check")["value"]
    pca = pd.read_csv(REPORTS / "forecast_pca_diagnostics.csv").iloc[0]
    importance = pd.read_csv(VOLUME / "xgboost_validation_feature_importance.csv")
    shap_values = pd.read_csv(VOLUME / "xgboost_validation_shap_summary.csv")
    mlflow_summary_path = REPORTS / "mlflow_run_summary.json"
    mlflow_summary = (
        json.loads(mlflow_summary_path.read_text(encoding="utf-8")) if mlflow_summary_path.is_file() else None
    )
    powerbi_spec_path = REPORTS / "powerbi" / "power_bi_dashboard_spec.json"
    powerbi_spec = (
        json.loads(powerbi_spec_path.read_text(encoding="utf-8")) if powerbi_spec_path.is_file() else None
    )
    horizon = int(manifest["forecast_horizon_days"])
    validation = metrics.loc[
        metrics["split"].eq("validation")
        & metrics["forecast_window_days"].eq(horizon)
    ].sort_values("MAE")
    test = metrics.loc[
        metrics["split"].eq("test")
        & metrics["forecast_window_days"].eq(horizon)
    ]
    winner = validation.iloc[0]
    winner_test = test.loc[test["model"].eq(winner["model"])].iloc[0]
    short_horizon = 7 if horizon != 7 else 30
    short_winner_name = manifest["selected_models_by_horizon"][str(short_horizon)]["selected_model"]
    short_test = metrics.loc[
        metrics["split"].eq("test")
        & metrics["forecast_window_days"].eq(short_horizon)
        & metrics["model"].eq(short_winner_name)
    ].iloc[0]
    priorities = sla.loc[~sla["priority"].eq("Overall")]
    priority_table = {
        "headers": ["Priority", "Reference", "Met / Assessed", "Attainment"],
        "rows": [
            [
                row.priority,
                f"{int(row.reference_hours)}h",
                f"{int(row.reference_met):,} / {int(row.assessed_tickets):,}",
                f"{row.reference_met_pct:.1f}%",
            ]
            for row in priorities.itertuples()
        ]
        + [
            [
                "Overall",
                "n/a",
                f"{int(sla.loc[sla['priority'].eq('Overall'), 'reference_met'].iloc[0]):,} / {int(sla.loc[sla['priority'].eq('Overall'), 'assessed_tickets'].iloc[0]):,}",
                f"{sla.loc[sla['priority'].eq('Overall'), 'reference_met_pct'].iloc[0]:.1f}%",
            ]
        ],
    }
    candidate_table = {
        "headers": ["Model", "Validation MAE", "Validation RMSE", "Test MAE"],
        "rows": [
            [
                row.model,
                f"{row.MAE:.2f}",
                f"{row.RMSE:.2f}",
                f"{test.loc[test['model'].eq(row.model), 'MAE'].iloc[0]:.2f}",
            ]
            for row in validation.itertuples()
        ],
    }
    if "forecast_window_days" in importance:
        importance = importance.loc[importance["forecast_window_days"].eq(horizon)]
    top_features = ", ".join(importance.head(3)["feature"].tolist())
    top_shap = ", ".join(shap_values.head(3)["feature"].tolist())
    overall_sla = sla.loc[sla["priority"].eq("Overall")].iloc[0]
    horizon_rows = sorted(
        [
            (short_horizon, short_winner_name, short_test),
            (horizon, winner["model"], winner_test),
        ],
        key=lambda item: item[0],
    )
    numeric_table = {
        "headers": ["Horizon", "Selected model", "Test MAE", "Test RMSE", "Test WAPE", "Test sMAPE", "90% coverage"],
        "rows": [
            [
                f"{days}-day",
                model,
                f"{row['MAE']:.2f}",
                f"{row['RMSE']:.2f}",
                f"{row['WAPE']:.1%}",
                f"{row['sMAPE']:.1%}",
                f"{row['coverage_90_interval']:.1%}",
            ]
            for days, model, row in horizon_rows
        ],
    }

    technical = [
        ("Help Desk Ticket SLA and Volume Forecasting", [
            "Technical capstone: retrospective resolution/SLA-reference analysis",
            "Core prediction: daily Ticket-arrival forecasts for the next week and month",
            "Historical experiment, not operational approval or current performance evidence",
        ]),
        ("Problem framing and measures", [
            "Retrospective: compare elapsed resolution time with configured priority references.",
            "Predictive: count exact Ticket-type issues created per UTC day; evaluate 7- and 30-day paths.",
            "Primary metrics: reference attainment by priority; forecast MAE/RMSE, WAPE/sMAPE and interval coverage.",
            "No live alert is defined. SLA references are inherited project assumptions, not verified policy.",
        ]),
        ("Source, scope and data quality", [
            "Identity: Kaggle v1 bytes and Mendeley v2 digest match. Original CC BY 4.0 conflicts with Kaggle's CC0 label.",
            f"{manifest['source_rows']:,} issue rows, {int(quality['source_columns'])} fields; exact Ticket type is the analysis cohort.",
            f"{manifest['scoped_ticket_rows']:,} scoped Ticket records, {manifest['date_start']} to {manifest['date_end']} UTC.",
            "Documentation states 2016 to March 2023, but the file contains pre-2016 timestamps; earlier rows are profiled and excluded.",
            f"{manifest['zero_ticket_days']:,}/{manifest['observed_calendar_days']:,} dates have no Ticket rows; treated as zero under unverified coverage assumption.",
        ], None, [
            "Monthly created-Ticket counts (notebook 02, EDA): volume varies across months and thins out in later years.",
        ], REPORTS / "notebook02_monthly_volume.png", 7.4, True),
        ("Retrospective reference results", [
            "Reference attainment by priority (exact Ticket type, 2016+, Done resolution, mapped priority, valid nonnegative duration):",
        ], priority_table, [
            "Blocker/Highest mapped to Critical 4h; High 8h; Medium 24h; Low/Lowest 24h; unknown excluded.",
            "Duration is UTC wall-clock time from issue_created to issue_resolution_date; business-hour calendars and pause rules are not verified.",
            "Results are reference comparisons, not contractual compliance or live warnings.",
        ]),
        ("Key numerical measures at a glance", [
            "Selected forecast model performance on the later, untouched historical test split, by horizon:",
        ], numeric_table, [
            f"Data scale: {manifest['source_rows']:,} source rows; {manifest['scoped_ticket_rows']:,} scoped Ticket records; {manifest['observed_calendar_days']:,} calendar days ({manifest['date_start']} to {manifest['date_end']}).",
            f"{manifest['zero_ticket_days']:,} dates ({manifest['zero_ticket_days'] / manifest['observed_calendar_days']:.1%}) have no Ticket rows, treated as zero under an unverified coverage assumption.",
            f"Retrospective reference attainment overall: {int(overall_sla['reference_met']):,} / {int(overall_sla['assessed_tickets']):,} assessed tickets ({overall_sla['reference_met_pct']:.1f}%).",
            f"Feature reduction: {int(pca['selected_features'])}/{int(pca['original_features'])} training-only MI-selected features; PCA used {int(pca['pca_components_for_target_variance'])} components for {pca['explained_variance_ratio']:.0%} of target-related variance.",
        ]),
        ("Process flow: source to decks", [
            "Eight notebooks, one canonical output/issue_helpdesk/ folder, then decks/report/app read from it:",
        ], None, [
            "Full step-by-step mapping with exact function and file names: docs/process_flow.md.",
        ], REPORTS / "process_flow_diagram.png"),
        ("Seven-day and month-ahead forecast design", [
            "Target: daily counts of exact Ticket records; UTC calendar.",
            "Separate 7- and 30-day expanding-window forecasts; the month horizon predicts daily counts, not one monthly total.",
            f"30-day experiment: {manifest['validation_forecast_origins_by_horizon']['30']} validation origins; {manifest['test_forecast_origins_by_horizon']['30']} later test origins.",
            "Compared seasonal naive, additive weekly ETS, XGBoost autoregression, and MI-selected/PCA Ridge.",
            "Validation MAE selects a model separately for each horizon; test remains separate.",
        ], None, [
            "30-day final-test MAE by lead-time band (notebook 03): error typically grows with lead time, one reason a 7-day path is tracked separately.",
        ], REPORTS / "notebook03_lead_time_mae.png", 7.4, True),
        ("Candidate comparison", [
            f"30-day validation winner: {winner['model']} (MAE {winner['MAE']:.2f} tickets/day).",
        ], candidate_table, [
            f"30-day selected-model test MAE {winner_test['MAE']:.2f}, RMSE {winner_test['RMSE']:.2f}, WAPE {winner_test['WAPE']:.1%}, sMAPE {winner_test['sMAPE']:.1%}.",
            f"7-day comparison: {short_winner_name}, test MAE {short_test['MAE']:.2f} tickets/day (its own validation-selected model).",
        ]),
        ("Uncertainty and explanation", [
            "MAE is the typical daily miss; RMSE emphasizes larger errors. WAPE/sMAPE express relative error.",
            f"Selected-model 90% test interval coverage: {winner_test['coverage_90_interval']:.1%}; calibrated on validation residuals.",
            "Panel: XGBoost SHAP, not ETS. Separate ETS state/residual evidence is linked in the report; fitted residuals are not test errors.",
        ], None, [
            f"Leading XGBoost validation features: {top_features}. Validation-only SHAP highlights: {top_shap}; importance is association, not causation. Feature selection kept {int(pca['selected_features'])}/{int(pca['original_features'])} MI-selected features.",
        ], REPORTS / "notebook04_explainability.png", 9.0, True),
        ("Ethics and fairness limitations", [
            "No approved protected demographic attributes: protected-group fairness cannot be assessed.",
            "Priority attainment uses priority-specific reference values; group rates need contextual interpretation.",
            "One aggregate forecast series cannot test allocation fairness across queue, site or shift.",
            "Small masked-project groups should not be ranked; retain human review and investigate persistent under-forecasting.",
            "Missing demographics, measured ROI or production are not automatic academic failures; explicit rubric requirements still apply.",
        ]),
        ("Deployment, MLOps and Power BI proof", [
            f"Selected model persisted and served via a FastAPI POC (src/app.py); MLflow run {mlflow_summary['run_id'][:12]} logs the same parameters and metrics for reproducibility."
            if mlflow_summary else
            "Selected model persisted and served via a FastAPI POC (src/app.py); MLflow tracks the same parameters and metrics for reproducibility.",
            f"MLflow experiment '{mlflow_summary['experiment_name']}': validation MAE {mlflow_summary['logged_metrics']['validation_mae']:.2f}, test MAE {mlflow_summary['logged_metrics']['test_mae']:.2f}."
            if mlflow_summary else
            "MLflow experiment tracking captures validation and test metrics for the selected run.",
            f"Power BI ready export folder ({powerbi_spec['report_fact']}, {len(powerbi_spec['pages'])} report pages, {len(powerbi_spec['screenshots']['files'])} rendered sample screenshots) covers both the after the fact SLA report and the volume forecast, rebuilt from source data on every run so it scales to a larger dataset."
            if powerbi_spec else
            "A Power BI ready export folder covers both the after the fact SLA report and the volume forecast, rebuilt from source data on every run.",
            "Power BI Desktop is not installed in this environment, so the screenshots are rendered stand ins from the same export tables, not native .pbix captures.",
            "Monitoring plan defines retrain triggers (schedule, drift, new source coverage) before any of this becomes an operational service.",
        ]),
        ("Integration readiness and conclusion", [
            "Ready locally: data profile, cohort rules, rolling backtest, validation selection, error/interval outputs and monitoring plan.",
            "Not ready: verified coverage, current arrivals, approved SLA rules, segmented forecasts or production integration.",
            "Next: confirm source range/zero days and policy; test on current future data against the seasonal baseline.",
            "Conclusion: stronger historical prototype, not an operationally validated staffing or compliance system.",
        ]),
    ]
    business = [
        ("Help Desk Ticket SLA and Volume Forecasting", [
            "Business-facing capstone summary",
            "Retrospective priority-reference insights and week/month daily-arrival forecasts",
            "Historical evidence to inform a possible human-reviewed capacity pilot",
        ]),
        ("The business questions", [
            "How often did completed Ticket records meet the configured priority references?",
            "Can the past daily arrival pattern support a useful next-week estimate?",
            "What evidence is still needed before planners use a forecast?",
        ]),
        ("What the data covers", [
            "Source identity verified. Retain Mendeley CC BY 4.0; resolve the conflicting Kaggle CC0 label.",
            f"{manifest['source_rows']:,} issues overall; {manifest['scoped_ticket_rows']:,} exact Ticket records in the documented date scope.",
            f"2,628 UTC dates from {manifest['date_start']} through {manifest['date_end']}.",
            f"{manifest['zero_ticket_days']:,} dates have no Ticket rows; the model assumes they are true zeros.",
            "The source file and documentation disagree about pre-2016 creation dates; current demand is not represented.",
        ]),
        ("Historical SLA-reference picture", [
            "Reference attainment by priority (see table below):",
        ], priority_table, [
            "Thresholds and priority mapping are carried forward as project settings, not owner-approved policy.",
            "Calendar wall-clock duration may not match the official SLA clock.",
        ]),
        ("Month-ahead forecast result", [
            f"30-day validation selected {winner['model']} with MAE {winner['MAE']:.2f} tickets/day.",
            f"Later historical test: MAE {winner_test['MAE']:.2f}, RMSE {winner_test['RMSE']:.2f}, WAPE {winner_test['WAPE']:.1%}.",
            f"Validation-calibrated 90% interval covered {winner_test['coverage_90_interval']:.1%} of test dates.",
            f"7-day comparison: {short_winner_name}, test MAE {short_test['MAE']:.2f} tickets/day.",
            "A month forecast is one estimate per day, not a guaranteed monthly total. These old results do not validate current demand.",
        ]),
        ("What the measures mean", [
            "MAE: average daily count miss. RMSE: larger misses weigh more.",
            "WAPE compares total misses to actual volume; sMAPE summarizes daily relative error.",
            "Interval coverage shows how often the historical test landed inside estimated bounds.",
            "Intervals and errors may change when source coverage or operations change.",
        ]),
        ("What remains unknown", [
            "Whether dates with no rows are fully covered zero-ticket days.",
            "Whether the historical, masked-project data matches the target service today.",
            "Whether the inherited priority targets and time clock match approved SLA policy.",
            "No staffing savings, avoided breach impact or ROI has been measured.",
        ]),
        ("Safe path forward", [
            "Resolve the mirror license label, source date-range mismatch and daily coverage.",
            "Obtain service-owner approval for priority mapping, references and SLA clock.",
            "Set forecast tolerances in advance and test on current, future-period data.",
            "Pilot with human review; compare service outcomes and planner effort before estimating value.",
        ]),
        ("Proof of work: a working prototype", [
            "This is not only analysis on paper: the selected forecast model is packaged and runnable today.",
            "A small web service (FastAPI) can return a forecast on request, showing the model works end to end.",
            "Every run is logged (model choice, accuracy numbers) so results can be traced and reproduced later.",
            "A ready to use Power BI folder holds both the SLA report and the volume forecast data, plus sample report screenshots, so a planner can see what the finished dashboard would look like.",
            "None of this is a live production system yet; it is evidence the idea is technically workable and reviewable.",
        ]),
        ("Recommendation", [
            "Use this as an improved historical analysis and forecast prototype.",
            "Do not automate SLA enforcement or staffing decisions from this evidence.",
            "Proceed only after data, policy and current-period validation.",
        ]),
    ]
    return {"technical_capstone.pptx": technical, "business_capstone.pptx": business}


def _add_bullets(text_frame, bullets, start_index=0, small=False):
    for offset, (text, size) in enumerate(zip(bullets, _bullet_sizes(bullets, small))):
        p = text_frame.paragraphs[0] if start_index == 0 and offset == 0 else text_frame.add_paragraph()
        p.text = f"\u2022  {text}"
        p.font.name = _font(size)[1]
        p.font.size = Pt(size)
        p.line_spacing = 1.25
        p.font.color.rgb = GRAY
        p.space_after = Pt(8 if small else 12)
    return start_index + len(bullets)


def _add_table(slide, table_spec, top):
    headers = table_spec["headers"]
    rows = table_spec["rows"]
    n_rows = len(rows) + 1
    n_cols = len(headers)
    row_heights = [
        max(
            Inches(0.42 + 0.5 / n_rows),
            max(_text_height([str(value)], Inches(11.5 / n_cols), [14], 0, index == 0)
                for value in row),
        )
        for index, row in enumerate([headers] + rows)
    ]
    height = sum(row_heights)
    shape = slide.shapes.add_table(n_rows, n_cols, Inches(0.9), top, Inches(11.5), height)
    table = shape.table
    for row, row_height in zip(table.rows, row_heights):
        row.height = row_height
    col_width = Inches(11.5 / n_cols)
    for col_index in range(n_cols):
        table.columns[col_index].width = col_width
    for col_index, header in enumerate(headers):
        cell = table.cell(0, col_index)
        cell.text = header
        cell.fill.solid()
        cell.fill.fore_color.rgb = NAVY
        run = cell.text_frame.paragraphs[0].runs[0]
        run.font.bold = True
        run.font.name = _font(14, True)[1]
        run.font.size = Pt(14)
        run.font.color.rgb = RGBColor(255, 255, 255)
    for row_index, row in enumerate(rows, start=1):
        band = RGBColor(255, 255, 255) if row_index % 2 else RGBColor(232, 240, 247)
        for col_index, value in enumerate(row):
            cell = table.cell(row_index, col_index)
            cell.text = str(value)
            cell.fill.solid()
            cell.fill.fore_color.rgb = band
            run = cell.text_frame.paragraphs[0].runs[0]
            run.font.size = Pt(14)
            run.font.name = _font(14)[1]
            run.font.color.rgb = NAVY if col_index == 0 else GRAY
    for row in table.rows:
        for cell in row.cells:
            cell.text_frame.word_wrap = True
            cell.text_frame.paragraphs[0].line_spacing = 1.25
    return height


def build_deck(filename, slides):
    """Reserve text and footer space before fitting proportional images; never truncate."""
    if not 8 <= len(slides) <= 12:
        raise ValueError(f"{filename} must have between 8 and 12 slides.")
    manifest = json.loads((VOLUME / "volume_forecast_experiment.json").read_text(encoding="utf-8"))
    pre_correction = manifest.get("forecast_calendar_policy") != FORECAST_CALENDAR_POLICY
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    total = len(slides)
    for number, entry in enumerate(slides, start=1):
        title, bullets = entry[0], entry[1]
        table_spec = entry[2] if len(entry) > 2 else None
        footnote = entry[3] if len(entry) > 3 else []
        image_path = entry[4] if len(entry) > 4 else None
        image_width = entry[5] if len(entry) > 5 else 11.0
        small_bullets = entry[6] if len(entry) > 6 else False
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        is_title_slide = number == 1
        context = f"{filename}, slide {number} ({title})"

        if is_title_slide:
            if table_spec or footnote or image_path:
                raise ValueError(f"{context}: title slide cannot include additional content.")
            if _text_height([title], Inches(11.5), [40], 0, True) > Inches(1.6):
                raise ValueError(f"{context}: title does not fit.")
            title_body_height = _text_height(bullets, Inches(11), [20] * len(bullets), 12)
            if title_body_height > CONTENT_BOTTOM - Inches(4.25):
                raise ValueError(f"{context}: title slide body does not fit.")
            slide.background.fill.solid()
            slide.background.fill.fore_color.rgb = NAVY
            heading = slide.shapes.add_textbox(Inches(0.9), Inches(2.55), Inches(11.5), Inches(1.6))
            heading.text_frame.word_wrap = True
            heading.text_frame.text = title
            paragraph = heading.text_frame.paragraphs[0]
            paragraph.font.name = _font(40, True)[1]
            paragraph.line_spacing = 1.25
            paragraph.font.size = Pt(40)
            paragraph.font.bold = True
            paragraph.font.color.rgb = RGBColor(255, 255, 255)
            accent = slide.shapes.add_shape(1, Inches(0.95), Inches(2.35), Inches(1.6), Inches(0.08))
            accent.fill.solid()
            accent.fill.fore_color.rgb = BLUE
            accent.line.fill.background()
            body = slide.shapes.add_textbox(Inches(0.95), Inches(4.25), Inches(11), title_body_height)
            body.text_frame.word_wrap = True
            for index, text in enumerate(bullets):
                p = body.text_frame.paragraphs[0] if index == 0 else body.text_frame.add_paragraph()
                p.text = text
                p.font.name = _font(20)[1]
                p.line_spacing = 1.25
                p.font.size = Pt(20)
                p.font.color.rgb = RGBColor(214, 226, 237)
                p.space_after = Pt(12)
        else:
            slide.background.fill.solid()
            slide.background.fill.fore_color.rgb = RGBColor(248, 250, 252)
            band = slide.shapes.add_shape(1, Inches(0), Inches(0), Inches(13.333), Inches(0.12))
            band.fill.solid()
            band.fill.fore_color.rgb = BLUE
            band.line.fill.background()
            heading = slide.shapes.add_textbox(Inches(0.65), Inches(0.42), Inches(12), Inches(0.8))
            heading.text_frame.word_wrap = True
            heading.text_frame.text = title
            paragraph = heading.text_frame.paragraphs[0]
            heading_size = 28 if len(title) < 48 else 24
            if _text_height([title], Inches(12), [heading_size], 0, True) > Inches(0.8):
                raise ValueError(f"{context}: heading does not fit.")
            paragraph.font.name = _font(heading_size, True)[1]
            paragraph.font.size = Pt(heading_size)
            paragraph.line_spacing = 1.25
            paragraph.font.bold = True
            paragraph.font.color.rgb = NAVY
            accent = slide.shapes.add_shape(1, Inches(0.68), Inches(1.32), Inches(1.25), Inches(0.07))
            accent.fill.solid()
            accent.fill.fore_color.rgb = BLUE
            accent.line.fill.background()
            body_top = Inches(1.68)
            note_height = _bullet_height(footnote, small=True)
            if bullets:
                body_height = _bullet_height(bullets, small=small_bullets)
                body = slide.shapes.add_textbox(Inches(0.9), body_top, Inches(11.5), body_height)
                body.text_frame.word_wrap = True
                _add_bullets(body.text_frame, bullets, small=small_bullets)
                body_top += body_height + CONTENT_GAP
            if table_spec:
                table_height = _add_table(slide, table_spec, body_top)
                body_top += table_height + CONTENT_GAP
            if image_path:
                if not image_path.is_file():
                    raise FileNotFoundError(f"{context}: missing image {image_path}")
                available = CONTENT_BOTTOM - body_top - note_height - (CONTENT_GAP if footnote else 0)
                if available < Inches(1):
                    raise ValueError(f"{context}: insufficient space for a readable image.")
                with Image.open(image_path) as image:
                    ratio = image.width / image.height
                width = min(Inches(image_width), Inches(11.5), int(available * ratio))
                height = int(width / ratio)
                if min(width, height) < Inches(1):
                    raise ValueError(f"{context}: image would be too small to read.")
                left = (prs.slide_width - width) // 2
                picture = slide.shapes.add_picture(str(image_path), left, body_top, width=width, height=height)
                body_top += picture.height + CONTENT_GAP
            if footnote:
                note = slide.shapes.add_textbox(Inches(0.9), body_top, Inches(11.5), note_height)
                note.text_frame.word_wrap = True
                _add_bullets(note.text_frame, footnote, small=True)
            if any(shape.top + shape.height > CONTENT_BOTTOM for shape in slide.shapes):
                raise ValueError(f"{context}: content does not fit above the footer at readable font sizes.")
            page_number = slide.shapes.add_textbox(Inches(12.2), Inches(7.08), Inches(0.9), Inches(0.28))
            page_number.text_frame.text = f"{number:02d} / {total:02d}"
            page_number.text_frame.paragraphs[0].font.size = Pt(10)
            page_number.text_frame.paragraphs[0].font.color.rgb = GRAY
        if pre_correction or not is_title_slide:
            footer = slide.shapes.add_textbox(Inches(0.65), Inches(7.08), Inches(9), Inches(0.28))
            footer.text_frame.text = (
                PRE_CORRECTION_WARNING if pre_correction else
                "Help Desk Ticket SLA and Volume Forecasting | Historical analysis"
            )
            paragraph = footer.text_frame.paragraphs[0]
            paragraph.font.name = _font(10, pre_correction)[1]
            paragraph.font.size = Pt(10)
            paragraph.font.bold = pre_correction
            paragraph.font.color.rgb = (
                RGBColor(255, 255, 255) if is_title_slide else
                RGBColor(153, 45, 25) if pre_correction else GRAY
            )
    OUTPUT.mkdir(parents=True, exist_ok=True)
    destination = OUTPUT / filename
    prs.save(destination)
    return destination


def build_all():
    paths = [build_deck(name, slides) for name, slides in deck_content().items()]
    for path in paths:
        print(f"Created {path}")
    return paths


if __name__ == "__main__":
    build_all()
